from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "SEO-AI-LEGAL-CONVERSION-AUDIT-2026-07-28.md"
OUTPUT = ROOT / "Accident-Claims-Scotland-Full-Audit-2026-07-29.docx"

BLUE = RGBColor(31, 77, 120)
DARK = RGBColor(11, 37, 69)
MUTED = RGBColor(90, 99, 110)
RED = RGBColor(155, 28, 28)
LIGHT = "E8EEF5"
GRAY = "F2F4F7"

UPDATE = """
# Revalidation update — 29 July 2026

This update supersedes any conflicting finding in the 28 July audit reproduced after it. The full strategic, keyword, content, AI-search, competitor, schema, internal-linking and answer-block recommendations remain applicable unless expressly updated below.

## Current release position

The codebase has materially improved since the original audit. Placeholder and unsupported regulatory-claim checks now pass; the enquiry form posts to a Cloudflare Worker that uses Resend; user-facing success and error states are truthful; the privacy notice names Ola Consultants Ltd as the recipient/controller; `LocalBusiness` schema has been replaced with `WebPage`; the unsupported `SearchAction` has been removed; a skip link is present; the sitemap now includes location and road-traffic detail pages; Cloudflare Pages and Worker configuration are committed; and lint and TypeScript checks pass.

The site is still not ready to present itself as a Scottish personal-injury solicitor service without further verification. The About page says the website is an information resource and is not a regulated law firm, while titles, introductions, location descriptions and calls to action repeatedly promise or imply “our specialist solicitors”, “accident claims solicitors serving [city]”, and no-win-no-fee handling. The exact relationship between Ola Consultants Ltd, the website, any receiving solicitor practice, the referral/lead model, service acceptance, fees, and data sharing remains insufficiently explicit.

## Revalidated top priorities

| Priority | Current action | Severity | Owner |
|---|---|---|---|
| 1 | Reconcile the business model sitewide. If this is an information/referral service, replace “our solicitors” and “we handle” with accurate wording; name every receiving provider before transfer and explain selection, responsibility and remuneration. | Critical | Compliance + content |
| 2 | Verify Ola Consultants Ltd’s corporate identity, registered details, ICO position, lawful bases and special-category-data condition; publish the verified facts and retention schedule. | Critical | Privacy/compliance |
| 3 | Complete end-to-end form testing in production: CORS, Worker secret, Resend delivery, mailbox monitoring, incident alerts, spam/rate limiting, response SLA and deletion/retention. | Critical | Developer + operations |
| 4 | Prevent enquiries from being accepted when delivery or monitoring is unavailable; add a reference number and accessible focus movement to the success panel. | High | Developer |
| 5 | Obtain page-by-page review from a named Scottish personal-injury solicitor and replace organisational authorship with real accountable authors/reviewers where appropriate. | High | Solicitor + content |
| 6 | Rewrite no-win-no-fee content against the actual provider’s written agreement. Remove categorical “no financial risk” language and explain success fees, VAT, outlays, insurance, adverse expenses and termination. | High | Solicitor/compliance |
| 7 | Noindex or consolidate city pages until each has distinct, evidenced local value; “serving” wording alone does not cure doorway-page risk. | High | SEO + content |
| 8 | Remove or merge same-intent guide/service duplicates identified in the URL map; implement 301 redirects and update internal links and sitemap entries. | High | SEO + developer |
| 9 | Replace blanket/future content dates with a verified publication and legal-review register. Update `lastmod` only after material review. | Medium | Content + developer |
| 10 | Establish measurement after lawful consent: Search Console, Bing Webmaster Tools, analytics, CRM/call outcomes and server/CDN logs. | High | Analytics + privacy |

## Confirmed changes since the original audit

| Original issue | Revalidated status | Evidence |
|---|---|---|
| Placeholder data on production | Fixed in code | `scripts/check-placeholders.mjs` passes; `siteConfig.ts` contains no placeholder contact or regulator data |
| Simulated form success | Fixed in code, production delivery still requires operational proof | `ClaimForm.tsx` posts JSON to the Worker and shows success only on an OK response |
| No server-side validation/abuse controls | Partly fixed | `workers/enquiry.js` validates origin, required fields and lengths, includes a honeypot and sends through Resend; rate limiting and operational alerting are not demonstrated |
| Unnamed data controller | Partly fixed | Privacy page names Ola Consultants Ltd, but corporate/ICO details and the full Article 9 basis need verification |
| Fabricated `LocalBusiness` schema | Fixed | Location pages now emit `WebPage` schema without local addresses |
| Unsupported `SearchAction` | Fixed | `WebSite` schema no longer declares site search |
| Missing skip link | Fixed | Root layout includes “Skip to main content” |
| Missing sitemap dates/routes | Partly fixed | Static, location, guide and road-traffic detail URLs have `lastModified`; dates still need a verified editorial basis |
| Lint errors | Fixed | ESLint and TypeScript completed successfully on 29 July 2026 |
| Build failure | Environment-limited, not a confirmed code defect | The production build reached Turbopack but the sandbox prevented a helper process binding to a port (`EPERM`) |
| Deployment unknown | Clarified in repository | `wrangler.toml` configures Cloudflare Pages static output; `wrangler.enquiry.toml` configures the enquiry Worker |
| Canonical host conflict | Code is internally consistent on `www`; live redirect/canonical behaviour still needs independent verification | `SITE.url`, `robots.txt`, sitemap and generated URLs use `https://www.accident-claims-scotland.com` |

## New or continuing confirmed issues

| Finding | Status | Why it matters | Exact fix |
|---|---|---|---|
| Identity and proposition conflict | Confirmed | About says information site/not a law firm; commercial templates imply the site’s own solicitors and legal service | Create one verified entity/service statement and apply it to titles, metadata, hero copy, forms, footer, legal pages and schema |
| No-win-no-fee wording is too categorical | Confirmed | “No financial risk” and “will not pay” statements may omit provider-specific exceptions, outlays, insurance and termination consequences | Use the actual written funding agreement; qualify every summary and identify the provider |
| Location pages remain highly templated | Confirmed | Doorway/thin-content risk persists despite safer schema | Noindex/merge until each page has genuine distinct usefulness and no false office implication |
| Organisational authorship without named legal reviewer | Confirmed | Weak YMYL accountability and AI trust | Add verified author/reviewer profiles, credentials, jurisdiction and real review dates |
| Metadata still describes “solicitors” as the site proposition | Confirmed | Search snippets can contradict the About disclosure | Rewrite root and route metadata to match the verified referral/information model |
| Form endpoint is cross-origin and publicly callable | Confirmed | Availability, abuse and privacy depend on the Worker configuration | Restrict allowed origins, rate-limit, log minimally, monitor delivery and publish processor/retention facts |
| Blanket sitemap `lastModified` date | Confirmed | Gives an unreliable freshness signal if pages were not substantively reviewed that day | Generate dates from a reviewed-content registry or source-control/content workflow |
| Search visibility not established | Confirmed from current public search sample | Exact-domain searches did not surface the target site in the sampled results | Submit/verify in GSC and Bing, inspect indexing/canonical selection, then prioritise content based on real query data |

## Live-search and legal-source note

Public search checks were repeated on 29 July 2026. The sampled results showed active Scottish competition from Winn Solicitors (Scotland), Scottish Accident, Lawford Kidd, Scottish Claims Helpline, Accident Claims Scotland `.co.uk`, Scotland Claims and other lead-generation/information sites. Exact searches for the target `.com` domain did not return the site in the sampled results; this is not proof of non-indexation and must be verified in Search Console and Bing Webmaster Tools.

Current legal review should continue to anchor limitation wording in sections 17–19A of the Prescription and Limitation (Scotland) Act 1973; funding in the Civil Litigation (Expenses and Group Proceedings) (Scotland) Act 2018 and the applicable success-fee regulations/agreement; CICA content in the current Criminal Injuries Compensation Scheme and official guidance; and procedural statements in current Scottish Courts and Tribunals Service rules and guidance. Nothing in this audit is individual legal advice.

## Release recommendation

Keep the site in an information-only, carefully qualified state until the provider/referral model and operational enquiry handling are independently verified. Do not activate solicitor-led claims language merely because form delivery works. The commercially strongest next move is to establish a transparent, named and accountable service chain, then align every page, schema object and conversion step with that reality.

"""


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clean_inline(text):
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 — \2", text)
    return text.strip()


def add_rich_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            p.add_run(re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 — \2", part))
    return p


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [clean_inline(c) for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    usable = 6.5
    widths = [usable / cols] * cols
    if cols == 4:
        widths = [0.55, 3.25, 1.15, 1.55]
    elif cols == 3:
        widths = [1.45, 2.65, 2.40]
    elif cols == 2:
        widths = [1.75, 4.75]
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.size = Pt(8.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = DARK
                shade(cell, LIGHT)
    set_repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown(doc, markdown):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and "|" in line[1:]:
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table(block))
            continue
        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(clean_inline(line[2:]), level=1)
        elif re.match(r"^\d+\.\s+", line):
            add_rich_paragraph(doc, re.sub(r"^\d+\.\s+", "", line), "List Number")
        elif line.startswith("- "):
            add_rich_paragraph(doc, line[2:], "List Bullet")
        elif line.startswith("> "):
            p = add_rich_paragraph(doc, line[2:])
            p.paragraph_format.left_indent = Inches(0.25)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = MUTED
        else:
            add_rich_paragraph(doc, line)
        i += 1


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, before, after, color in [
    ("Heading 1", 16, 16, 8, BLUE),
    ("Heading 2", 13, 12, 6, BLUE),
    ("Heading 3", 11.5, 8, 4, DARK),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
for name in ("List Bullet", "List Number"):
    styles[name].font.name = "Calibri"
    styles[name].font.size = Pt(10.5)
    styles[name].paragraph_format.space_after = Pt(4)
    styles[name].paragraph_format.line_spacing = 1.10

header = section.header.paragraphs[0]
header.text = "ACCIDENT CLAIMS SCOTLAND  |  INDEPENDENT AUDIT"
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in header.runs:
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = MUTED

footer = section.footer.paragraphs[0]
footer.text = "SEO • AI search • Scots-law content • trust • conversion  |  29 July 2026"
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ACCIDENT CLAIMS SCOTLAND")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RED
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Full SEO, AI Search,\nLegal Content and Conversion Audit")
r.bold = True
r.font.size = Pt(28)
r.font.color.rgb = DARK
p.paragraph_format.space_after = Pt(16)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Revalidated 29 July 2026")
r.font.size = Pt(14)
r.font.color.rgb = BLUE
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Codebase, deployment configuration, public-search sample and authoritative-source review")
r.font.size = Pt(10.5)
r.font.color.rgb = MUTED
doc.add_page_break()

add_markdown(doc, UPDATE)
doc.add_page_break()

source_text = SOURCE.read_text(encoding="utf-8")
source_text = source_text.replace(
    "# Accident Claims Scotland: SEO, AI Search, Legal Content and Conversion Audit",
    "# Original full audit — 28 July 2026",
    1,
)
add_markdown(doc, source_text)

doc.core_properties.title = "Accident Claims Scotland — Full SEO, AI Search, Legal Content and Conversion Audit"
doc.core_properties.subject = "Revalidated website and codebase audit"
doc.core_properties.author = "Independent audit prepared with Codex"
doc.core_properties.keywords = "SEO, AI search, Scottish personal injury, legal content, conversion, accessibility"
doc.save(OUTPUT)
print(OUTPUT)
